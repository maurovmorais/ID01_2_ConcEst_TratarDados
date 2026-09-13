"""Leitura de arquivos Excel utilizando openpyxl.

Este módulo concentra a classe responsável por abrir um arquivo .xlsx e
devolver os dados de suas abas em estruturas Python simples (lista de
dicionários), para serem tratados/gravados por outras camadas da
automação (ex.: leitura_depara.py -> manipular_tabelas.py).
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook

logger = logging.getLogger(__name__)


class LeitorExcel:
    """Encapsula a abertura e leitura de um arquivo .xlsx com openpyxl.

    Uso recomendado (garante o fechamento do arquivo mesmo em caso de
    erro):

        with LeitorExcel(caminho) as leitor:
            linhas = leitor.ler_planilha("Empresas")
    """

    def __init__(self, caminho_arquivo: str | Path) -> None:
        self.caminho_arquivo = Path(caminho_arquivo)
        self._workbook: Optional[Workbook] = None

    def __enter__(self) -> "LeitorExcel":
        self.abrir()
        return self

    def __exit__(self, exc_type, exc_value, traceback) -> None:
        self.fechar()

    def abrir(self) -> None:
        """Abre o arquivo Excel em modo somente leitura."""
        if not self.caminho_arquivo.exists():
            raise FileNotFoundError(
                f"Arquivo Excel não encontrado: {self.caminho_arquivo}"
            )
        logger.debug("Abrindo arquivo Excel: %s", self.caminho_arquivo)
        self._workbook = load_workbook(
            filename=self.caminho_arquivo, data_only=True, read_only=True
        )

    def fechar(self) -> None:
        """Fecha o workbook, liberando o arquivo."""
        if self._workbook is not None:
            self._workbook.close()
            self._workbook = None

    @property
    def workbook(self) -> Workbook:
        if self._workbook is None:
            raise RuntimeError(
                "Workbook não está aberto. Use 'with LeitorExcel(...) as "
                "leitor:' ou chame 'abrir()' antes de ler os dados."
            )
        return self._workbook

    def listar_planilhas(self) -> list[str]:
        """Retorna os nomes de todas as abas do arquivo."""
        return list(self.workbook.sheetnames)

    def ler_planilha(
        self,
        nome_planilha: str,
        primeira_linha_cabecalho: bool = True,
    ) -> list[dict[str, object]]:
        """Lê todas as linhas de dados de uma planilha (aba).

        Args:
            nome_planilha: nome exato da aba dentro do arquivo Excel.
            primeira_linha_cabecalho: se True (padrão), usa a primeira
                linha da aba como nomes das colunas (chaves do
                dicionário retornado).

        Returns:
            Lista de dicionários, um por linha de dados. Linhas
            totalmente vazias (todas as células None) são ignoradas.

        Raises:
            ValueError: se a planilha informada não existir no arquivo.
        """
        if nome_planilha not in self.workbook.sheetnames:
            raise ValueError(
                f"Planilha '{nome_planilha}' não encontrada no arquivo "
                f"'{self.caminho_arquivo.name}'. "
                f"Disponíveis: {self.listar_planilhas()}"
            )

        planilha = self.workbook[nome_planilha]
        linhas = planilha.iter_rows(values_only=True)

        colunas: Optional[list[str]] = None
        if primeira_linha_cabecalho:
            cabecalho = next(linhas, ())
            colunas = [
                str(valor).strip() if valor is not None else f"coluna_{indice}"
                for indice, valor in enumerate(cabecalho)
            ]

        dados: list[dict[str, object]] = []
        for linha in linhas:
            if all(valor is None for valor in linha):
                continue
            if colunas is None:
                colunas = [f"coluna_{indice}" for indice in range(len(linha))]
            dados.append(dict(zip(colunas, linha)))

        logger.info(
            "Planilha '%s' lida com sucesso: %d linha(s) de dados.",
            nome_planilha,
            len(dados),
        )
        return dados
